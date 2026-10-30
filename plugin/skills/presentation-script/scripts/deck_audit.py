#!/usr/bin/env python3
"""발표 덱 점검(읽기 전용). 사용: python3 deck_audit.py 덱.pptx [강조색 HEX=E0000F]
장별로 (1) 위쪽 글줄(주장 문장 후보) (2) 강조색이 들어간 **도형(상자·글상자) 수**를 센다(한 상자 안의 여러 글자는 1곳).
강조색 상자가 4곳 넘는 장은 ⚠로 표시한다(구분 슬라이드·표지는 장식이라 무시해도 된다). python-pptx 필요: uv run --with python-pptx python deck_audit.py ...
"""
import sys
from pptx import Presentation

def _has_color(r):
    try: return r.font.color.rgb is not None
    except Exception: return False

def main():
    if len(sys.argv) < 2: sys.exit(__doc__)
    accent = (sys.argv[2] if len(sys.argv) > 2 else "E0000F").upper()
    p = Presentation(sys.argv[1])
    print(f"{sys.argv[1]} — {len(p.slides)}장, 강조색 {accent}\n")
    for n, s in enumerate(p.slides, 1):
        runs = fills = 0
        label = ""
        for sh in s.shapes:
            try:
                if sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == accent: fills += 1
            except Exception: pass
            if sh.has_text_frame:
                if not label and sh.is_placeholder and sh.text_frame.text.strip(): label = sh.text_frame.text.strip().split("\n")[0]
                if any(str(r.font.color.rgb) == accent and r.text.strip() and r.text.strip() != "."
                       for para in sh.text_frame.paragraphs for r in para.runs if _has_color(r)):
                    runs += 1
        total = runs + fills
        flag = "⚠ 강조색 많음" if total > 4 else ""
        sentence = "문장형" if len(label) >= 14 else "짧은 제목"
        print(f"{n:>2}장 | 강조 글상자 {runs:>2} · 도형 {fills:>2} | {sentence} | {label[:44]} {flag}")

main()
