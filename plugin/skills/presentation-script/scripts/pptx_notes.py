#!/usr/bin/env python3
"""PPT 발표자 메모 도구 (python-pptx 필요). 슬라이드 본문은 건드리지 않는다.
  dump 덱.pptx                       슬라이드별 제목·본문·메모를 출력
  add  덱.pptx 메모.json [--force]    메모.json = [{"match":"제목 앞부분","note":"메모 전문"}, ...]
                                     제목으로 슬라이드를 찾는다(번호 사용 안 함). 백업 후 저장.
"""
import sys, os, json, copy, shutil, argparse
from pptx import Presentation

def titles(slide):
    out = []
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6: walk(sh.shapes)
            elif sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text.split("\n")[0])
    walk(slide.shapes); return out

def body_frame(slide, ref):
    ns = slide.notes_slide
    if ns.notes_text_frame is None:   # 메모 칸이 없는 장은 다른 장의 본문 칸을 복제한다
        b = [s for s in ref.notes_slide.shapes if s.is_placeholder and s.placeholder_format.type == 2][0]
        ns.shapes._spTree.append(copy.deepcopy(b._element))
    return ns.notes_text_frame

def dump(path):
    p = Presentation(path); print(f"{len(p.slides)}장")
    for i, s in enumerate(p.slides, 1):
        n = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame is not None else ""
        print(f"\n===== {i} =====\n" + "\n".join(titles(s)[:8]) + (f"\n[NOTES] {n}" if n.strip() else ""))

def add(path, notes_json, force):
    d, f = os.path.split(os.path.abspath(path))
    lock = os.path.join(d, "~$" + f)
    if os.path.exists(lock) and not force:
        sys.exit(f"⚠ 잠금 파일이 있습니다({lock}) — PowerPoint가 열고 있을 수 있습니다. 닫은 뒤 다시 하거나 --force.")
    items = json.load(open(notes_json, encoding="utf-8"))
    p = Presentation(path); slides = list(p.slides)
    ref = next((s for s in slides if s.has_notes_slide and s.notes_slide.notes_text_frame is not None), slides[0])
    plan = []
    for it in items:
        hit = [i for i, s in enumerate(slides) if any(t.startswith(it["match"]) for t in titles(s)[:3])]
        if len(hit) != 1: sys.exit(f"'{it['match']}'에 맞는 슬라이드가 {len(hit)}개입니다 — 제목을 더 구체적으로.")
        plan.append((hit[0], it["note"]))
    bak = path + ".bak"; shutil.copy2(path, bak)
    for i, note in plan: body_frame(slides[i], ref).text = note
    tmp = path + ".tmp"; p.save(tmp); os.replace(tmp, path)
    print(f"{len(plan)}장에 메모를 넣었습니다. 백업: {bak}")

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); sp = ap.add_subparsers(dest="cmd", required=True)
    d = sp.add_parser("dump"); d.add_argument("deck")
    a = sp.add_parser("add"); a.add_argument("deck"); a.add_argument("notes"); a.add_argument("--force", action="store_true")
    args = ap.parse_args()
    dump(args.deck) if args.cmd == "dump" else add(args.deck, args.notes, args.force)
